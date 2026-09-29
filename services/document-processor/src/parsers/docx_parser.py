"""DOCX parser for Word documents."""

import os
from typing import Dict, Any
from datetime import datetime, timezone
from docx import Document
from loguru import logger

from parsers.base_parser import BaseParser
from models import ParsedDocument


class DocxParser(BaseParser):
    """Parser for Microsoft Word documents."""
    
    def __init__(self):
        super().__init__()
        self.supported_extensions = {'.docx', 'docx'}
    
    async def parse(
        self, 
        file_path: str,
        extract_tables: bool = True,
        **kwargs
    ) -> ParsedDocument:
        """
        Parse DOCX document.
        
        Args:
            file_path: Path to DOCX file
            extract_tables: Whether to extract tables
            
        Returns:
            ParsedDocument with extracted content
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")
        
        start_time = datetime.now(timezone.utc)
        file_size = os.path.getsize(file_path)
        
        try:
            logger.info(f"Starting DOCX parsing ({file_size:,} bytes)")
            
            # Open document
            doc = Document(file_path)
            
            # Extract paragraphs
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text.strip())
            
            # Extract tables
            tables_text = []
            table_count = 0
            if extract_tables:
                for table in doc.tables:
                    table_count += 1
                    table_data = []
                    for row in table.rows:
                        row_data = [cell.text.strip() for cell in row.cells]
                        if any(row_data):
                            table_data.append(" | ".join(row_data))
                    
                    if table_data:
                        tables_text.append(f"\n[TABLE {table_count}]\n" + "\n".join(table_data) + "\n[/TABLE]\n")
            
            # Combine content
            content = "\n\n".join(paragraphs)
            if tables_text:
                content += "\n\n" + "\n".join(tables_text)
            
            # Extract metadata
            metadata = {
                "paragraphs": len(paragraphs),
                "tables": table_count,
                "file_size": file_size,
            }
            
            # Try to extract document properties
            try:
                core_properties = doc.core_properties
                if core_properties:
                    metadata.update({
                        "title": core_properties.title or "",
                        "author": core_properties.author or "",
                        "subject": core_properties.subject or "",
                        "keywords": core_properties.keywords or "",
                        "created": core_properties.created.isoformat() if core_properties.created else "",
                        "modified": core_properties.modified.isoformat() if core_properties.modified else "",
                    })
            except Exception as e:
                logger.warning(f"Could not extract DOCX metadata: {e}")
            
            processing_time = (datetime.now(timezone.utc) - start_time).total_seconds()
            
            logger.info(
                f"DOCX parsing complete: {len(paragraphs)} paragraphs, "
                f"{table_count} tables, {len(content):,} chars in {processing_time:.2f}s"
            )
            
            return ParsedDocument(
                content=content,
                metadata=metadata,
                processing_time=processing_time
            )
            
        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
            raise